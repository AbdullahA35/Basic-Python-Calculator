"""Run every case in test_cases.xlsx and write Actual Result and Status back.

Usage (from the project root):  python scripts/update_excel_results.py
"""
import importlib
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
WORKBOOK = ROOT / "test_cases.xlsx"


def run_case(function_path, a, b, expected):
    module_name, function_name = function_path.split(".")
    function = getattr(importlib.import_module(f"calculator.{module_name}"), function_name)
    try:
        actual = function(a, b)
    except Exception as error:  # the expected value may be an exception name
        name = type(error).__name__
        return name, name == expected
    if expected.endswith("Error"):
        return repr(actual), False
    return repr(actual), abs(actual - float(expected)) < 1e-9


def main():
    workbook = openpyxl.load_workbook(WORKBOOK)
    sheet = workbook["Test Cases"]
    header = {cell.value: cell.column for cell in sheet[1]}
    passed = 0
    for row in range(2, sheet.max_row + 1):
        if not sheet.cell(row, 1).value:
            continue
        function_path = sheet.cell(row, header["Module / Function"]).value
        a = sheet.cell(row, header["Input A"]).value
        b = sheet.cell(row, header["Input B"]).value
        expected = str(sheet.cell(row, header["Expected Result"]).value)
        actual, ok = run_case(function_path, a, b, expected)
        sheet.cell(row, header["Actual Result"]).value = actual
        sheet.cell(row, header["Status"]).value = "Pass" if ok else "Fail"
        passed += ok
    workbook.save(WORKBOOK)
    print(f"Updated {WORKBOOK.name}: {passed} passed")


if __name__ == "__main__":
    main()
